from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from BE.app.db import SessionLocal, init_db
from BE.app.models import Gym

PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = PROJECT_ROOT / "BE" / "Temp" / "kolkata_25_real_gyms_updated.xlsx"

REQUIRED_HEADERS = {"name", "address", "latitude", "longitude"}


def load_rows(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"Gym workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Gyms"] if "Gyms" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ValueError(f"Workbook has no rows: {path}")
        headers = [str(value).strip().lower() if value is not None else "" for value in header_row]
        missing = REQUIRED_HEADERS - set(headers)
        if missing:
            raise ValueError(f"Workbook is missing required columns: {', '.join(sorted(missing))}")
        records = []
        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            record = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
            if not str(record.get("name") or "").strip():
                raise ValueError(f"Gym name is missing at spreadsheet row {row_number}")
            records.append(record)
        return records
    finally:
        workbook.close()


def _as_text(value: Any) -> str | None:
    if value is None or not str(value).strip():
        return None
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%H:%M")
        except ValueError:
            pass
    return str(value).strip()


GYM_TEXT_FIELDS = (
    "about", "status", "address", "city", "state", "pincode", "google_maps_url",
    "phone", "email", "image_url", "monday_open", "monday_close", "tuesday_open",
    "tuesday_close", "wednesday_open", "wednesday_close", "thursday_open", "thursday_close",
    "friday_open", "friday_close", "saturday_open", "saturday_close", "sunday_open",
    "sunday_close", "equipment_json", "photos_json",
)


def sync_gyms_to_db(db: Session, rows: list[dict[str, Any]] | None = None) -> tuple[int, int]:
    rows = rows if rows is not None else load_rows(WORKBOOK_PATH)
    inserted = 0
    updated = 0
    seen_names: set[str] = set()

    for row in rows:
        name = str(row["name"]).strip()
        key = name.casefold()
        if key in seen_names:
            raise ValueError(f"Duplicate gym name in workbook: {name}")
        seen_names.add(key)

        gym = db.query(Gym).filter(Gym.name.ilike(name)).first()
        if gym is None:
            gym = Gym(name=name, status="active")
            db.add(gym)
            inserted += 1
        else:
            updated += 1

        for field in GYM_TEXT_FIELDS:
            value = _as_text(row.get(field))
            if value is not None:
                setattr(gym, field, value)

        for field in ("latitude", "longitude"):
            value = row.get(field)
            if value is not None and str(value).strip():
                setattr(gym, field, float(value))

        for field, default in (("slot_duration_minutes", 30), ("max_bookings_per_slot", 1)):
            value = row.get(field)
            if value is not None and str(value).strip():
                setattr(gym, field, int(float(value)))
            elif getattr(gym, field) is None:
                setattr(gym, field, default)

    db.commit()
    return inserted, updated


def main() -> None:
    init_db()
    rows = load_rows(WORKBOOK_PATH)
    print(f"Read {len(rows)} gym rows from {WORKBOOK_PATH}")

    db = SessionLocal()
    try:
        inserted, updated = sync_gyms_to_db(db, rows)
        print(
            f"Imported {inserted} new gyms and updated {updated} existing gyms "
            f"({len(rows)} workbook rows)."
        )
    finally:
        db.close()


if __name__ == "__main__":
    main()
